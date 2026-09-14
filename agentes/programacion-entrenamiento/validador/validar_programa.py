"""Validador determinista (Paso 3) del Asistente de Programación de Entrenamiento.

Comprueba mecánicamente, contra un archivo `.xlsx` en el formato descrito en
`formato-salida/formato-excel.md`, todo lo que NO requiere juicio (eso lo hace
el Crítico, Paso 4). Esto no es una descripción de las comprobaciones — es el
código que las ejecuta, para que dejen de leerse a ojo (ver `docs/roadmap.md`,
backlog histórico).

Uso:
    python3 validar_programa.py <programa.xlsx> \
        [--catalogo catalogo-ejercicios.xlsx] \
        [--excluidos "Hack Squat" "Elevación lateral en máquina"] \
        [--semanas-esperadas 3]

Salida: informe JSON por stdout con `aprobado`, `errores` y `advertencias`,
y código de salida 0 si aprobado, 1 si no.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field

import openpyxl

HOJA_PROGRAMA = "Programa"
HOJA_PROGRAMACION = "Programación"

COLUMNAS_PROGRAMACION = [
    "semana", "dia", "nombre_dia", "es_descanso", "notas_dia",
    "bloque", "instrucciones_bloque", "ejercicio", "equipo",
    "series", "reps", "rir", "rpe", "carga_kg", "carga_pct",
    "descanso_seg", "tempo", "duracion_seg", "notas",
]

COLUMNAS_EJERCICIO = COLUMNAS_PROGRAMACION[7:]  # de 'ejercicio' a 'notas' (col. 8 a 19)


@dataclass
class InformeValidacion:
    errores: list[str] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)

    @property
    def aprobado(self) -> bool:
        return not self.errores

    def to_dict(self) -> dict:
        return {
            "aprobado": self.aprobado,
            "errores": self.errores,
            "advertencias": self.advertencias,
        }


def _normalizar(texto: str | None) -> str:
    if not texto:
        return ""
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sin_acentos).strip().lower()


def _es_verdadero(valor) -> bool:
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, str):
        return valor.strip().upper() == "TRUE"
    return False


def _celda_vacia(valor) -> bool:
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


def cargar_catalogo(ruta_catalogo: str) -> set[str]:
    wb = openpyxl.load_workbook(ruta_catalogo, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return set()
    cabecera = [_normalizar(c) for c in filas[0]]
    try:
        idx_titulo = cabecera.index("title")
    except ValueError:
        try:
            idx_titulo = cabecera.index("titulo")
        except ValueError:
            idx_titulo = len(filas[0]) - 1  # última columna como fallback razonable
    return {_normalizar(fila[idx_titulo]) for fila in filas[1:] if fila and not _celda_vacia(fila[idx_titulo])}


def validar_programa(
    ruta_programa: str,
    ruta_catalogo: str | None = None,
    ejercicios_excluidos: list[str] | None = None,
    semanas_esperadas: int | None = None,
) -> InformeValidacion:
    informe = InformeValidacion()
    excluidos_normalizados = {_normalizar(e) for e in (ejercicios_excluidos or [])}
    catalogo = cargar_catalogo(ruta_catalogo) if ruta_catalogo else None

    wb = openpyxl.load_workbook(ruta_programa, data_only=True)

    # 1. Nombres de hoja exactos.
    if HOJA_PROGRAMA not in wb.sheetnames:
        informe.errores.append(f"Falta la hoja obligatoria '{HOJA_PROGRAMA}'.")
    if HOJA_PROGRAMACION not in wb.sheetnames:
        informe.errores.append(f"Falta la hoja obligatoria '{HOJA_PROGRAMACION}'.")
    if informe.errores:
        return informe  # sin las dos hojas no se puede seguir comprobando nada más

    ws_programa = wb[HOJA_PROGRAMA]
    ws_prog = wb[HOJA_PROGRAMACION]

    # 2. Cabecera de 'Programa' y valor de 'semanas'.
    cab_programa = [c.value for c in ws_programa[1]]
    fila_programa = [c.value for c in ws_programa[2]] if ws_programa.max_row >= 2 else []
    semanas_declaradas = None
    if "semanas" in cab_programa:
        idx = cab_programa.index("semanas")
        if idx < len(fila_programa):
            semanas_declaradas = fila_programa[idx]
    else:
        informe.advertencias.append("La hoja 'Programa' no tiene columna 'semanas' (es informativa, no bloqueante).")

    # 3. Cabecera de 'Programación': columnas exactas presentes (orden como advertencia).
    cab_prog = [c.value for c in ws_prog[1]]
    faltantes = [c for c in COLUMNAS_PROGRAMACION if c not in cab_prog]
    if faltantes:
        informe.errores.append(f"Faltan columnas obligatorias en '{HOJA_PROGRAMACION}': {faltantes}.")
        return informe  # sin las columnas no se puede indexar el resto de filas con fiabilidad
    if cab_prog[: len(COLUMNAS_PROGRAMACION)] != COLUMNAS_PROGRAMACION:
        informe.advertencias.append(
            "Las columnas de 'Programación' no están en el orden estándar del formato (no bloqueante)."
        )
    idx_col = {nombre: cab_prog.index(nombre) for nombre in COLUMNAS_PROGRAMACION}

    def val(fila, nombre):
        return fila[idx_col[nombre]] if idx_col[nombre] < len(fila) else None

    filas = [f for f in ws_prog.iter_rows(min_row=2, values_only=True) if any(c is not None for c in f)]
    if not filas:
        informe.errores.append(f"La hoja '{HOJA_PROGRAMACION}' no tiene ninguna fila de datos.")
        return informe

    semanas_vistas: set[int] = set()
    dias_por_semana: dict[int, dict[int, str]] = {}

    for n_fila, fila in enumerate(filas, start=2):
        semana = val(fila, "semana")
        dia = val(fila, "dia")
        nombre_dia = val(fila, "nombre_dia")
        es_descanso = _es_verdadero(val(fila, "es_descanso"))
        ejercicio = val(fila, "ejercicio")
        series = val(fila, "series")
        reps = val(fila, "reps")
        rir = val(fila, "rir")
        rpe = val(fila, "rpe")

        # 4. 'semana' y 'dia' obligatorias y con rango válido.
        if _celda_vacia(semana):
            informe.errores.append(f"Fila {n_fila}: falta 'semana'.")
        elif not isinstance(semana, (int, float)) or int(semana) != semana or semana < 1:
            informe.errores.append(f"Fila {n_fila}: 'semana' debe ser un entero positivo (valor: {semana!r}).")
        else:
            semanas_vistas.add(int(semana))

        if _celda_vacia(dia):
            informe.errores.append(f"Fila {n_fila}: falta 'dia'.")
        elif not isinstance(dia, (int, float)) or int(dia) != dia or not (1 <= dia <= 7):
            informe.errores.append(f"Fila {n_fila}: 'dia' debe ser un entero entre 1 y 7 (valor: {dia!r}).")

        if _celda_vacia(nombre_dia):
            informe.errores.append(f"Fila {n_fila}: falta 'nombre_dia'.")
        elif isinstance(semana, (int, float)) and isinstance(dia, (int, float)):
            clave_semana = int(semana)
            clave_dia = int(dia)
            existentes = dias_por_semana.setdefault(clave_semana, {})
            if clave_dia in existentes and existentes[clave_dia] != nombre_dia:
                informe.errores.append(
                    f"Fila {n_fila}: 'nombre_dia' ({nombre_dia!r}) no coincide con el ya visto "
                    f"para semana {clave_semana}, día {clave_dia} ({existentes[clave_dia]!r})."
                )
            else:
                existentes[clave_dia] = nombre_dia

        if es_descanso:
            # 5. Fila de descanso: columnas de ejercicio (8-19) deben quedar vacías.
            rellenas = [c for c in COLUMNAS_EJERCICIO if not _celda_vacia(val(fila, c))]
            if rellenas:
                informe.errores.append(
                    f"Fila {n_fila}: es_descanso=TRUE pero tiene columnas de ejercicio rellenas: {rellenas}."
                )
            continue

        # 6. Fila de ejercicio: 'ejercicio', 'series' y 'reps' obligatorias.
        if _celda_vacia(ejercicio):
            informe.errores.append(f"Fila {n_fila}: falta 'ejercicio' (obligatorio salvo fila de descanso).")
        if _celda_vacia(series):
            informe.errores.append(f"Fila {n_fila}: falta 'series' (obligatorio salvo fila de descanso).")
        if _celda_vacia(reps) and _celda_vacia(val(fila, "duracion_seg")):
            informe.errores.append(
                f"Fila {n_fila}: falta 'reps' (obligatorio salvo fila de descanso o ejercicio por tiempo con 'duracion_seg')."
            )

        # 7. rir y rpe simultáneos: advertencia, no error (formato-excel.md: "no es un error, pero solo se usa rpe").
        if not _celda_vacia(rir) and not _celda_vacia(rpe):
            informe.advertencias.append(
                f"Fila {n_fila}: 'rir' y 'rpe' rellenos a la vez ({rir!r} / {rpe!r}) — se usará 'rpe', "
                "considera dejar solo uno."
            )

        # 8. Ejercicios excluidos para este cliente (lesión, material no disponible...).
        if not _celda_vacia(ejercicio) and _normalizar(ejercicio) in excluidos_normalizados:
            informe.errores.append(
                f"Fila {n_fila}: el ejercicio {ejercicio!r} está en la lista de exclusiones de este cliente."
            )

        # 9. Coincidencia con el catálogo real (si se ha proporcionado) — advertencia, no error:
        #    un ejercicio nuevo no es un fallo, ver 'Frontera de responsabilidad' en formato-excel.md.
        if catalogo is not None and not _celda_vacia(ejercicio) and _normalizar(ejercicio) not in catalogo:
            informe.advertencias.append(
                f"Fila {n_fila}: el ejercicio {ejercicio!r} no coincide exactamente con ningún título del catálogo "
                "— se creará como ejercicio nuevo si el matcher de BeFit no encuentra una coincidencia aproximada."
            )

    # 10. 'semanas' de la hoja Programa coincide con las semanas distintas realmente escritas.
    if semanas_declaradas is not None and semanas_vistas:
        if int(semanas_declaradas) != len(semanas_vistas):
            informe.errores.append(
                f"La hoja 'Programa' declara semanas={semanas_declaradas!r} pero 'Programación' contiene "
                f"{len(semanas_vistas)} semana(s) distinta(s): {sorted(semanas_vistas)}."
            )

    # 11. Todas las semanas pedidas están explícitas (sin huecos), si se indica cuántas se esperaban.
    if semanas_esperadas is not None:
        esperado = set(range(1, semanas_esperadas + 1))
        faltan = esperado - semanas_vistas
        sobran = semanas_vistas - esperado
        if faltan:
            informe.errores.append(f"Faltan semanas explícitas: {sorted(faltan)} (se esperaban 1..{semanas_esperadas}).")
        if sobran:
            informe.advertencias.append(f"Hay semanas no esperadas en el archivo: {sorted(sobran)}.")
    elif semanas_vistas and semanas_vistas != set(range(1, max(semanas_vistas) + 1)):
        informe.errores.append(
            f"Las semanas escritas tienen huecos: {sorted(semanas_vistas)} (deberían ser consecutivas desde 1)."
        )

    return informe


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("programa", help="Ruta al archivo .xlsx del programa a validar.")
    parser.add_argument("--catalogo", default=None, help="Ruta a catalogo-ejercicios.xlsx (opcional).")
    parser.add_argument(
        "--excluidos", nargs="*", default=None,
        help="Nombres de ejercicios excluidos para este cliente (lesión, material no disponible...).",
    )
    parser.add_argument("--semanas-esperadas", type=int, default=None, help="Número de semanas que el programa debería cubrir.")
    args = parser.parse_args()

    informe = validar_programa(
        args.programa,
        ruta_catalogo=args.catalogo,
        ejercicios_excluidos=args.excluidos,
        semanas_esperadas=args.semanas_esperadas,
    )
    print(json.dumps(informe.to_dict(), ensure_ascii=False, indent=2))
    return 0 if informe.aprobado else 1


if __name__ == "__main__":
    sys.exit(main())
