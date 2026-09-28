"""Pruebas del validador de progresión de macrociclo (`validar_macrociclo.py`).

A diferencia de `test_validar_programa.py` (que usa como fixture el programa real de un cliente), este
validador comprueba una propiedad MATEMÁTICA del conjunto de mesociclos -- si las series suben, si el pico no
retrocede, si las reps cambian de rango -- independiente del cliente o de sus ejercicios concretos. No hay
ninguna razón de individualización para usar un caso real como fixture aquí (ver la nota de procedencia del
propio validador), así que las pruebas construyen un macrociclo sintético mínimo con `openpyxl`, correcto por
diseño, y lo mutan deliberadamente para provocar cada fallo uno a uno -- mismo espíritu de
`test_validar_programa.py`, con fixtures generadas en vez de un archivo real adjunto.
"""

import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from validar_macrociclo import validar_macrociclo  # noqa: E402

HEAD = ["semana", "dia", "nombre_dia", "es_descanso", "notas_dia", "bloque", "instrucciones_bloque", "ejercicio",
        "equipo", "series", "reps", "rir", "rpe", "carga_kg", "carga_pct", "descanso_seg", "tempo",
        "duracion_seg", "notas", "tecnica", "tecnica_series"]

# (nombre, patrón A=descendente / B=ascendente, series base en M1) -- 4 ejercicios, 2 sesiones, para tener
# ambos sentidos representados en cada sesión (requisito de V3).
EJERCICIOS = [
    ("Press banca barra", 1, "A", 3),
    ("Curl biceps mancuerna", 1, "B", 2),
    ("Sentadilla barra", 2, "A", 3),
    ("Elevacion lateral mancuerna", 2, "B", 2),
]
REPS_A = ["12-15", "10-12", "8-10", "12-15"]  # S1, S2, S3, S4(deload=S1)
REPS_B = ["6-8", "8-10", "10-12", "6-8"]


def _build_mesociclo(ruta, numero_mesociclo, series_extra=0, semanas=4, forzar_reps_fijas=False):
    """Mesociclo de `semanas` semanas (última = deload), con 4 ejercicios en 2 días. `series_extra` se suma a
    la base de M1 de cada ejercicio para simular que el mesociclo N tiene más volumen que el N-1 -- pásalo
    igual en todos los mesociclos de una prueba para simular un macrociclo PLANO (debe fallar V1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programa"
    ws.append(["titulo", "descripcion", "semanas"])
    ws.append([f"Mesociclo {numero_mesociclo}", f"Mesociclo {numero_mesociclo} de prueba", semanas])
    wp = wb.create_sheet("Programación")
    wp.append(HEAD)
    for w in range(1, semanas + 1):
        es_deload = w == semanas
        for nombre, dia, patron, base in EJERCICIOS:
            base_meso = base + series_extra
            if es_deload:
                series = max(1, base_meso // 2)
            elif 1 < w < semanas:
                series = base_meso + 1  # onda: semanas intermedias +1 serie (patrón real de progresion-carga.md)
            else:
                series = base_meso
            if forzar_reps_fijas:
                reps = "10-12"
            else:
                tabla = REPS_A if patron == "A" else REPS_B
                reps = tabla[0] if es_deload else tabla[min(w - 1, 2)]
            wp.append([w, dia, f"Dia {dia}", "FALSE", None, "Parte principal", None, nombre, "Barra",
                       series, reps, "2", None, None, None, 120, None, None, None, None, None])
    wb.save(ruta)


def _build_mesociclo_explicito(ruta, numero_mesociclo, s1, onda, semanas=4):
    """Mesociclo con una sola serie de ejercicios y una onda explícita: S1 = `s1`, S2 = S3 = `s1 + onda`,
    deload = mitad de S1. Solo para el test de V2b, donde hace falta control exacto del pico, no de un
    incremento uniforme entre mesociclos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programa"
    ws.append(["titulo", "descripcion", "semanas"])
    ws.append([f"Mesociclo {numero_mesociclo}", f"Mesociclo {numero_mesociclo} de prueba", semanas])
    wp = wb.create_sheet("Programación")
    wp.append(HEAD)
    for w in range(1, semanas + 1):
        es_deload = w == semanas
        series = max(1, s1 // 2) if es_deload else (s1 + onda if 1 < w < semanas else s1)
        for nombre, dia, patron, _base in EJERCICIOS:
            tabla = REPS_A if patron == "A" else REPS_B
            reps = tabla[0] if es_deload else tabla[min(w - 1, 2)]
            wp.append([w, dia, f"Dia {dia}", "FALSE", None, "Parte principal", None, nombre, "Barra",
                       series, reps, "2", None, None, None, 120, None, None, None, None, None])
    wb.save(ruta)


class TestMacrocicloSintetico(unittest.TestCase):
    """Macrociclo de 4 mesociclos, cada uno con +1 serie base y volumen dentro de rango respecto al anterior:
    debe aprobar sin errores (caso base correcto por construcción)."""

    def setUp(self):
        self.tmp = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_tmp_macrociclo")
        os.makedirs(self.tmp, exist_ok=True)
        self.rutas = []
        for i in range(1, 5):
            ruta = os.path.join(self.tmp, f"M{i}.xlsx")
            _build_mesociclo(ruta, i, series_extra=i - 1)
            self.rutas.append(ruta)

    def tearDown(self):
        for r in self.rutas:
            os.remove(r)
        os.rmdir(self.tmp)

    def test_aprueba_macrociclo_con_progresion_correcta(self):
        informe = validar_macrociclo(self.rutas)
        self.assertEqual(informe.errores, [])
        self.assertTrue(informe.aprobado)

    def test_un_solo_mesociclo_no_es_error_solo_advertencia(self):
        informe = validar_macrociclo([self.rutas[0]])
        self.assertTrue(informe.aprobado)
        self.assertTrue(any("menos de 2" in a for a in informe.advertencias))

    def test_detecta_volumen_plano_entre_mesociclos(self):
        # Reconstruye 3 mesociclos IDÉNTICOS (series_extra=0 en todos) -- V1 debe fallar para cada grupo.
        rutas_planas = []
        for i in range(1, 4):
            ruta = os.path.join(self.tmp, f"plano_M{i}.xlsx")
            _build_mesociclo(ruta, i, series_extra=0)
            rutas_planas.append(ruta)
        try:
            informe = validar_macrociclo(rutas_planas)
            self.assertFalse(informe.aprobado)
            self.assertTrue(any(e.startswith("V1 ") for e in informe.errores))
        finally:
            for r in rutas_planas:
                os.remove(r)

    def test_detecta_pico_que_retrocede_aunque_la_base_suba(self):
        # Reproduce el caso real (Josu, ver CHANGELOG.md v0.24.0): M1 (ya en curso) aplica una onda agresiva
        # (+2 en semanas intermedias) a todos los ejercicios; M2 sube la base (S1) pero su onda se diseñó más
        # conservadora (+0) -- el pico de M2 queda por DEBAJO del pico de M1 aunque la base de M2 sea mayor.
        ruta_m1 = os.path.join(self.tmp, "pico_M1.xlsx")
        ruta_m2 = os.path.join(self.tmp, "pico_M2.xlsx")
        _build_mesociclo_explicito(ruta_m1, 1, s1=6, onda=2)  # S1=6, S2=S3=8 (pico 8)
        _build_mesociclo_explicito(ruta_m2, 2, s1=7, onda=0)  # S1=7 (> 6, V1 ok), S2=S3=7 (pico 7 < 8)
        try:
            informe = validar_macrociclo([ruta_m1, ruta_m2])
            self.assertFalse(informe.aprobado)
            self.assertTrue(any(e.startswith("V2b ") for e in informe.errores), informe.errores)
        finally:
            os.remove(ruta_m1)
            os.remove(ruta_m2)

    def test_detecta_reps_fijas_sin_progresion(self):
        ruta = os.path.join(self.tmp, "reps_fijas.xlsx")
        _build_mesociclo(ruta, 1, forzar_reps_fijas=True)
        ruta2 = os.path.join(self.tmp, "reps_fijas_M2.xlsx")
        _build_mesociclo(ruta2, 2, series_extra=1, forzar_reps_fijas=True)
        try:
            informe = validar_macrociclo([ruta, ruta2])
            self.assertFalse(informe.aprobado)
            self.assertTrue(any(e.startswith("V3 ") for e in informe.errores))
        finally:
            os.remove(ruta)
            os.remove(ruta2)

    def test_agrupar_hombro_fusiona_grupos(self):
        informe_sin = validar_macrociclo(self.rutas, agrupar_hombro=False)
        informe_con = validar_macrociclo(self.rutas, agrupar_hombro=True)
        tabla_con = "\n".join(t for t in informe_con.tablas if t.startswith("== V1"))
        self.assertIn("Hombros:", tabla_con)
        self.assertTrue(informe_sin.aprobado and informe_con.aprobado)

    def test_patron_seguro_detecta_violacion_de_minimos(self):
        # "Press banca barra" con RIR "2" en todas las semanas de carga cumple RIR>=2; exige RIR>=3 para forzar el fallo.
        informe = validar_macrociclo(self.rutas, patron_seguro="press banca", rir_min_seguro=3)
        self.assertFalse(informe.aprobado)
        self.assertTrue(any(e.startswith("V4") for e in informe.errores))

    def test_orden_declarado_no_coincide_con_posicion_da_advertencia(self):
        # Renombra el título del segundo archivo para que declare "Mesociclo 5" en vez de 2.
        wb = openpyxl.load_workbook(self.rutas[1])
        wb["Programa"]["A2"] = "Mesociclo 5"
        ruta = os.path.join(self.tmp, "orden_raro.xlsx")
        wb.save(ruta)
        try:
            informe = validar_macrociclo([self.rutas[0], ruta] + self.rutas[2:])
            self.assertTrue(any("por orden de archivo ocupa la posición" in a for a in informe.advertencias))
        finally:
            os.remove(ruta)


if __name__ == "__main__":
    unittest.main()
