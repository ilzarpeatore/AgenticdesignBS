"""Lector compartido del `.xlsx` de un mesociclo (formato descrito en
`../formato-salida/formato-excel.md`).

Extraído de `validar_programa.py` (v0.6.0) para que `agentes/control-producto/`
(que necesita leer "lo prescrito" para compararlo contra la ejecución real,
ver su system-prompt.md sección 3, Paso 3) interprete el mismo formato de
columnas exactamente igual que el validador -- en vez de que cada uno lo
parsee por su cuenta y diverjan con el tiempo si el formato cambia.

Este módulo solo lee y valida la forma del archivo (hojas y columnas
correctas). No aplica ninguna regla de negocio (eso sigue siendo de
`validar_programa.py` para lo mecánico, y del propio Control Producto para
su comparación contra datos reales).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass

import openpyxl

HOJA_PROGRAMA = "Programa"
HOJA_PROGRAMACION = "Programación"

COLUMNAS_PROGRAMACION = [
    "semana", "dia", "nombre_dia", "es_descanso", "notas_dia",
    "bloque", "instrucciones_bloque", "ejercicio", "equipo",
    "series", "reps", "rir", "rpe", "carga_kg", "carga_pct",
    "descanso_seg", "tempo", "duracion_seg", "notas",
]

COLUMNAS_EJERCICIO = COLUMNAS_PROGRAMACION[7:]  # de 'ejercicio' a 'notas' (col. 8 a 19)


def normalizar(texto: str | None) -> str:
    if not texto:
        return ""
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sin_acentos).strip().lower()


def es_verdadero(valor) -> bool:
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, str):
        return valor.strip().upper() == "TRUE"
    return False


def celda_vacia(valor) -> bool:
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


class LecturaProgramaError(Exception):
    """El archivo no tiene la forma mínima para poder leerse con fiabilidad
    (hoja obligatoria ausente, columnas obligatorias ausentes, o sin ninguna
    fila de datos). `errores` trae el mismo texto que antes emitía
    `validar_programa.py` en estos casos -- el llamador solo tiene que
    añadirlo a su propio informe, no reformularlo."""

    def __init__(self, errores: list[str]):
        super().__init__("; ".join(errores))
        self.errores = errores


@dataclass
class LecturaPrograma:
    filas: list[tuple]
    idx_col: dict[str, int]
    orden_estandar: bool
    tiene_columna_semanas: bool
    semanas_declaradas: object | None

    def val(self, fila, nombre):
        idx = self.idx_col[nombre]
        return fila[idx] if idx < len(fila) else None


def leer_programa(ruta_programa: str) -> LecturaPrograma:
    wb = openpyxl.load_workbook(ruta_programa, data_only=True)

    errores_forma: list[str] = []
    if HOJA_PROGRAMA not in wb.sheetnames:
        errores_forma.append(f"Falta la hoja obligatoria '{HOJA_PROGRAMA}'.")
    if HOJA_PROGRAMACION not in wb.sheetnames:
        errores_forma.append(f"Falta la hoja obligatoria '{HOJA_PROGRAMACION}'.")
    if errores_forma:
        raise LecturaProgramaError(errores_forma)

    ws_programa = wb[HOJA_PROGRAMA]
    ws_prog = wb[HOJA_PROGRAMACION]

    cab_programa = [c.value for c in ws_programa[1]]
    fila_programa = [c.value for c in ws_programa[2]] if ws_programa.max_row >= 2 else []
    semanas_declaradas = None
    tiene_columna_semanas = "semanas" in cab_programa
    if tiene_columna_semanas:
        idx = cab_programa.index("semanas")
        if idx < len(fila_programa):
            semanas_declaradas = fila_programa[idx]

    cab_prog = [c.value for c in ws_prog[1]]
    faltantes = [c for c in COLUMNAS_PROGRAMACION if c not in cab_prog]
    if faltantes:
        raise LecturaProgramaError(
            [f"Faltan columnas obligatorias en '{HOJA_PROGRAMACION}': {faltantes}."]
        )
    orden_estandar = cab_prog[: len(COLUMNAS_PROGRAMACION)] == COLUMNAS_PROGRAMACION
    idx_col = {nombre: cab_prog.index(nombre) for nombre in COLUMNAS_PROGRAMACION}

    filas = [f for f in ws_prog.iter_rows(min_row=2, values_only=True) if any(c is not None for c in f)]
    if not filas:
        raise LecturaProgramaError([f"La hoja '{HOJA_PROGRAMACION}' no tiene ninguna fila de datos."])

    return LecturaPrograma(
        filas=filas,
        idx_col=idx_col,
        orden_estandar=orden_estandar,
        tiene_columna_semanas=tiene_columna_semanas,
        semanas_declaradas=semanas_declaradas,
    )
